import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

SECTPR_ORDER = [
    'headerReference', 'footerReference', 'type', 'pgSz', 'pgMar', 'paperSrc',
    'pgBorders', 'lnNumType', 'pgNumType', 'cols', 'formProt', 'vAlign',
    'noEndnote', 'titlePg', 'textDirection', 'bidi', 'rtlGutter', 'docGrid', 'printerSettings'
]

TBLPR_ORDER = [
    'tblStyle', 'tblpPr', 'tblOverlap', 'bidiVisual', 'tblStyleRowBandSize',
    'tblStyleColBandSize', 'tblW', 'jc', 'tblCellSpacing', 'tblInd', 'tblBorders',
    'shd', 'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption', 'tblDescription', 'tblPrChange'
]

TCPR_ORDER = [
    'tcW', 'gridSpan', 'hMerge', 'vMerge', 'tcBorders', 'shd', 'noWrap',
    'tcMar', 'textDirection', 'tcFitText', 'vAlign', 'hideMark', 'headers',
    'cellIns', 'cellDel', 'cellMerge', 'tcPrChange'
]

def reorder_children(parent, order_list):
    children = list(parent)
    def sort_key(elem):
        tag = elem.tag.split('}')[-1]
        try:
            return order_list.index(tag)
        except ValueError:
            return 999
    sorted_children = sorted(children, key=sort_key)
    for c in children:
        parent.remove(c)
    for c in sorted_children:
        parent.append(c)

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=80, bottom=80, left=120, right=120):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E0", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_double_page_borders(section):
    sectPr = section._sectPr
    pgBorders = parse_xml(f'''
        <w:pgBorders {nsdecls("w")} w:offsetFrom="page">
            <w:top w:val="double" w:sz="18" w:space="24" w:color="1A365D"/>
            <w:left w:val="double" w:sz="18" w:space="24" w:color="1A365D"/>
            <w:bottom w:val="double" w:sz="18" w:space="24" w:color="1A365D"/>
            <w:right w:val="double" w:sz="18" w:space="24" w:color="1A365D"/>
        </w:pgBorders>
    ''')
    sectPr.append(pgBorders)

def add_page_number(run):
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def build_docx_workbook():
    from workbook_data import weeks_data
    docx_path = "Fenix_IT_Mall_Project_Work_Book.docx"
    print(f"Building {docx_path}...")
    doc = Document()

    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)

    add_double_page_borders(section)

    header = section.header
    hp = header.paragraphs[0]
    hp.text = "GEETANJALI COLLEGE • SAURASHTRA UNIVERSITY        |        PROJECT WORK BOOK (2025-2026)"
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in hp.runs:
        r.font.name = "Times New Roman"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r_foot = fp.add_run("Candidate: Jeet Dodiya | Project: Fenix IT Mall                                              Page ")
    r_foot.font.name = "Times New Roman"
    r_foot.font.size = Pt(8.5)
    r_foot.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
    frun = fp.add_run()
    frun.font.name = "Times New Roman"
    frun.font.size = Pt(8.5)
    add_page_number(frun)

    # Color Constants
    c_blue = RGBColor(0x1A, 0x36, 0x5D)
    c_hand = RGBColor(0x1E, 0x3A, 0x8A) # Dark blue ink
    c_green = RGBColor(0x05, 0x96, 0x69)

    # Helper function for student handwritten paragraphs
    def add_hand_run(p, text, is_bold=False):
        r = p.add_run(text)
        r.font.name = "Segoe Print"
        r.font.size = Pt(9.5)
        r.bold = is_bold
        r.font.color.rgb = c_hand
        return r

    # =========================================================================
    # PAGE 1: COVER PAGE
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("GEETANJALI COLLEGE OF COMPUTER SCIENCE & COMMERCE")
    r.bold = True
    r.font.size = Pt(14)
    r.font.color.rgb = c_blue

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("Affiliated to Saurashtra University, Rajkot")
    r.font.size = Pt(11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run("DEPARTMENT OF COMPUTER SCIENCE")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("STUDENT PROJECT INTERNSHIP WORK BOOK")
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = c_blue

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(18)
    r = p.add_run("WEEKLY DEVELOPMENT LOG & CONTINUOUS EVALUATION DIARY")
    r.bold = True
    r.font.size = Pt(10.5)
    r.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("PROJECT TITLE :")
    r.bold = True
    r.font.size = Pt(11)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("Fenix IT Mall")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(22)
    r = p.add_run("Smart Retail Inventory & Real-Time Point of Sale Management System")
    r.italic = True
    r.font.size = Pt(11)

    meta_items = [
        ("Candidate Name :", "Jeet Dodiya", True),
        ("Course / Semester :", "B.C.A. / M.Sc. (IT) – Semester V / VI", False),
        ("Academic Year :", "2025 – 2026", False),
        ("Training Period :", "15/06/2026 to 15/09/2026 (3 Months / 14 Weeks)", True),
        ("Development Stack :", "Python 3.10+, Django 5.0, SQLite/PostgreSQL, Vanilla CSS3, Alpine.js", False),
        ("Internal Project Guide :", "Prof. Harsh Joshi / Prof. Kishorsinh Vala", False),
        ("Head of Department :", "Prof. Brijesh Shah", False),
        ("Institution :", "Geetanjali College of Computer Science and Commerce, Rajkot", False),
    ]
    t_cov = doc.add_table(rows=len(meta_items), cols=2)
    t_cov.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_cov, color="1A365D")
    c_widths = [Inches(2.4), Inches(4.2)]
    for ri, (lbl, val, is_b) in enumerate(meta_items):
        c1 = t_cov.cell(ri, 0)
        c2 = t_cov.cell(ri, 1)
        c1.width = c_widths[0]
        c2.width = c_widths[1]
        set_cell_background(c1, "F8FAFC")
        set_cell_background(c2, "F8FAFC")
        set_cell_margins(c1, top=80, bottom=80, left=120, right=120)
        set_cell_margins(c2, top=80, bottom=80, left=120, right=120)
        
        cp1 = c1.paragraphs[0]
        r1 = cp1.add_run(lbl)
        r1.bold = True
        r1.font.size = Pt(10)
        
        cp2 = c2.paragraphs[0]
        r2 = cp2.add_run(val)
        r2.bold = is_b
        r2.font.size = Pt(10)
        if is_b:
            r2.font.color.rgb = c_blue

    doc.add_page_break()

    # =========================================================================
    # PAGE 2: STUDENT UNDERTAKING
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("STUDENT DECLARATION & UNDERTAKING")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = c_blue

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("I, ")
    r2 = p.add_run("Jeet Dodiya")
    r2.bold = True
    r3 = p.add_run(", student of Geetanjali College of Computer Science and Commerce, hereby solemnly declare that the project entitled ")
    r4 = p.add_run("\"Fenix IT Mall\"")
    r4.bold = True
    r5 = p.add_run(" is a bona fide record of original development work carried out by me personally with my own hands during the 3-month period from ")
    r6 = p.add_run("15th June 2026 to 15th September 2026")
    r6.bold = True
    r7 = p.add_run(" under the guidance of ")
    r8 = p.add_run("Prof. Harsh Joshi")
    r8.bold = True
    r9 = p.add_run(".")

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.line_spacing = Pt(16)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run("I explicitly declare that all database models, entity schemas, application logic in Django, custom Vanilla CSS styling, responsive layouts, POS terminal calculation algorithms, barcode label generation, and the complete suite of 195 automated unit test cases were designed, coded, tested, and debugged by my own hands. No third-party agency or outsourced developer was hired for this work.")

    t_box = doc.add_table(rows=1, cols=1)
    t_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_box, color="3B82F6")
    cell = t_box.cell(0, 0)
    set_cell_background(cell, "EFF6FF")
    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
    cp = cell.paragraphs[0]
    cp.paragraph_format.line_spacing = Pt(15)
    r1 = cp.add_run("Statement Regarding Learning & Reference Tools:\n")
    r1.bold = True
    r1.font.color.rgb = c_blue
    r2 = cp.add_run("During the course of independent research, standard technical documentation (Official Django Documentation, MDN Web Docs, Python Standard Library) and AI study reference tools (such as Google Gemini) were utilized strictly as an interactive technical reference to clear conceptual doubts, understand architectural patterns, and resolve syntactical queries. No code was copied blindly or unverified. Every algorithm and feature was typed, integrated, executed, and thoroughly verified by me on my local machine.")
    r2.font.size = Pt(9.5)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(28)

    t_sigs = doc.add_table(rows=2, cols=2)
    t_sigs.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_widths = [Inches(3.3), Inches(3.3)]
    sig_data = [
        ("Date: 15/09/2026\nPlace: Rajkot", "\nCandidate Signature:\n\n_______________________\nJeet Dodiya"),
        ("\n\nInternal Guide Signature:\n\n_______________________\nProf. Harsh Joshi", "\n\nHead of Department:\n\n_______________________\nProf. Brijesh Shah")
    ]
    for ri, (txt1, txt2) in enumerate(sig_data):
        c1 = t_sigs.cell(ri, 0)
        c2 = t_sigs.cell(ri, 1)
        c1.width = s_widths[0]
        c2.width = s_widths[1]
        c1.paragraphs[0].text = txt1
        c2.paragraphs[0].text = txt2
        c1.paragraphs[0].paragraph_format.space_after = Pt(6)
        c2.paragraphs[0].paragraph_format.space_after = Pt(6)
        if ri == 0:
            c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            c2.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

    # =========================================================================
    # PAGE 3: 14-WEEK AT-A-GLANCE MATRIX
    # =========================================================================
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("14-WEEK INTERNSHIP DEVELOPMENT SCHEDULE")
    r.bold = True
    r.font.size = Pt(16)
    r.font.color.rgb = c_blue

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run("Continuous Evaluation & Milestone Tracking (15/06/2026 to 15/09/2026)")
    r.font.size = Pt(10)
    r.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

    matrix_rows = [
        ["Wk", "Date Range", "Milestone / Module", "Status", "Guide Init."],
        ["1", "15/06 – 21/06", "Problem Identification, Store Visits & Scope Definition", "Completed", "H.J."],
        ["2", "22/06 – 28/06", "SRS Documentation, Hardware/Software Feasibility", "Completed", "H.J."],
        ["3", "29/06 – 05/07", "Entity-Relationship (ER) Modeling & 3NF Normalization", "Completed", "H.J."],
        ["4", "06/07 – 12/07", "Data Flow Diagrams (Level 0, 1 & 2) & Wireframing", "Completed", "H.J."],
        ["5", "13/07 – 19/07", "Django 5.0 Project Setup, Virtualenv & Git Repository", "Completed", "H.J."],
        ["6", "20/07 – 26/07", "Custom User Model, RBAC Roles & Profile Security", "Completed", "H.J."],
        ["7", "27/07 – 02/08", "Product Catalog, Auto Barcode (EAN-13) & Duplicate Check", "Completed", "H.J."],
        ["8", "03/08 – 09/08", "Responsive Glassmorphic UI, Dashboard & Chart.js", "Completed", "H.J."],
        ["9", "10/08 – 16/08", "POS Terminal Core, Barcode Scanner & Real-Time GST", "Completed", "H.J."],
        ["10", "17/08 – 23/08", "Purchase Orders, Supplier Ledger & Reconciliations", "Completed", "H.J."],
        ["11", "24/08 – 30/08", "Staff Daily Attendance, Leave Approval & CRM Module", "Completed", "H.J."],
        ["12", "31/08 – 06/09", "Sequential Invoice Series, ReportLab PDF Invoicing", "Completed", "H.J."],
        ["13", "07/09 – 13/09", "195 Automated Unit Tests & Quality Assurance Audit", "Completed", "H.J."],
        ["14", "14/09 – 15/09", "Cloud Deployment on Render, Project Report & Viva Prep", "Completed", "H.J."],
    ]
    t_mat = doc.add_table(rows=len(matrix_rows), cols=5)
    t_mat.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(t_mat, color="CBD5E0")
    m_widths = [Inches(0.4), Inches(1.3), Inches(3.6), Inches(0.9), Inches(0.5)]
    for ri, row in enumerate(matrix_rows):
        for ci, val in enumerate(row):
            cell = t_mat.cell(ri, ci)
            cell.width = m_widths[ci]
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            if ri == 0:
                set_cell_background(cell, "1A365D")
            else:
                set_cell_background(cell, "F8FAFC" if ri % 2 == 1 else "FFFFFF")
            cp = cell.paragraphs[0]
            cp.paragraph_format.space_before = Pt(2)
            cp.paragraph_format.space_after = Pt(2)
            if ci in [0, 1, 3, 4]:
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cp.add_run(val)
            run.font.size = Pt(8.5)
            if ri == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            elif ci == 3:
                run.bold = True
                run.font.color.rgb = c_green

    doc.add_page_break()

    # =========================================================================
    # DETAILED WEEKS (1 TO 14)
    # =========================================================================
    for w in weeks_data:
        # Header Box
        t_wh = doc.add_table(rows=1, cols=1)
        t_wh.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t_wh, color="3182CE")
        cell = t_wh.cell(0, 0)
        set_cell_background(cell, "EBF8FF")
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        cp = cell.paragraphs[0]
        r1 = cp.add_run(f"WEEK {w['week']} : {w['module'].upper()}\n")
        r1.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = c_blue
        r2 = cp.add_run(f"Timeline: {w['dates']} • Focus: {w['planned']}")
        r2.font.size = Pt(9)
        r2.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

        p_sp = doc.add_paragraph()
        p_sp.paragraph_format.space_before = Pt(4)
        p_sp.paragraph_format.space_after = Pt(2)

        # Daily logs table
        t_dl = doc.add_table(rows=len(w['daily']) + 1, cols=2)
        t_dl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t_dl, color="CBD5E0")
        dl_widths = [Inches(0.9), Inches(5.7)]

        # Header row
        c0 = t_dl.cell(0, 0)
        c1 = t_dl.cell(0, 1)
        c0.width = dl_widths[0]
        c1.width = dl_widths[1]
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "F1F5F9")
        set_cell_margins(c0, top=60, bottom=60, left=80, right=80)
        set_cell_margins(c1, top=60, bottom=60, left=80, right=80)
        r = c0.paragraphs[0].add_run("Date")
        r.bold = True
        r.font.size = Pt(9.5)
        r = c1.paragraphs[0].add_run("Student's Daily Work & Development Activities (In Student's Own Hand)")
        r.bold = True
        r.font.size = Pt(9.5)

        for ri, (dt, log) in enumerate(w['daily']):
            cell_dt = t_dl.cell(ri + 1, 0)
            cell_log = t_dl.cell(ri + 1, 1)
            cell_dt.width = dl_widths[0]
            cell_log.width = dl_widths[1]
            set_cell_background(cell_dt, "FFFFFF" if ri % 2 == 0 else "F8FAFC")
            set_cell_background(cell_log, "FFFFFF" if ri % 2 == 0 else "F8FAFC")
            set_cell_margins(cell_dt, top=50, bottom=50, left=80, right=80)
            set_cell_margins(cell_log, top=50, bottom=50, left=80, right=80)
            
            cp_dt = cell_dt.paragraphs[0]
            r_dt = cp_dt.add_run(dt)
            r_dt.bold = True
            r_dt.font.size = Pt(9)
            r_dt.font.color.rgb = c_blue

            cp_log = cell_log.paragraphs[0]
            add_hand_run(cp_log, f"• {log}")

        p_sp2 = doc.add_paragraph()
        p_sp2.paragraph_format.space_before = Pt(4)
        p_sp2.paragraph_format.space_after = Pt(2)

        # Reflection table
        refl_items = [
            ("Self-Study & Learnings:", w["learnings"]),
            ("Engineering Challenges Solved:", w["challenges"]),
            ("Key Deliverables Produced:", w["deliverable"])
        ]
        t_rf = doc.add_table(rows=3, cols=2)
        t_rf.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t_rf, color="CBD5E0")
        rf_widths = [Inches(2.1), Inches(4.5)]
        for ri, (lbl, val) in enumerate(refl_items):
            c1 = t_rf.cell(ri, 0)
            c2 = t_rf.cell(ri, 1)
            c1.width = rf_widths[0]
            c2.width = rf_widths[1]
            set_cell_background(c1, "FAFAFA")
            set_cell_background(c2, "FAFAFA")
            set_cell_margins(c1, top=60, bottom=60, left=80, right=80)
            set_cell_margins(c2, top=60, bottom=60, left=80, right=80)
            
            cp1 = c1.paragraphs[0]
            r1 = cp1.add_run(lbl)
            r1.bold = True
            r1.font.size = Pt(9)

            cp2 = c2.paragraphs[0]
            if ri < 2:
                add_hand_run(cp2, val)
            else:
                r2 = cp2.add_run(val)
                r2.bold = True
                r2.font.size = Pt(9.5)

        p_sp3 = doc.add_paragraph()
        p_sp3.paragraph_format.space_before = Pt(4)
        p_sp3.paragraph_format.space_after = Pt(2)

        # Guide Remarks Box
        t_gb = doc.add_table(rows=1, cols=2)
        t_gb.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(t_gb, color="10B981")
        c1 = t_gb.cell(0, 0)
        c2 = t_gb.cell(0, 1)
        c1.width = Inches(5.1)
        c2.width = Inches(1.5)
        set_cell_background(c1, "ECFDF5")
        set_cell_background(c2, "ECFDF5")
        set_cell_margins(c1, top=80, bottom=80, left=100, right=100)
        set_cell_margins(c2, top=80, bottom=80, left=100, right=100)

        cp1 = c1.paragraphs[0]
        r1 = cp1.add_run("Guide Remarks : ")
        r1.bold = True
        r1.font.size = Pt(9)
        r2 = cp1.add_run(f"\"{w['guide_remarks']}\"")
        r2.italic = True
        r2.font.size = Pt(9)
        r2.font.color.rgb = RGBColor(0x06, 0x5F, 0x46)

        cp2 = c2.paragraphs[0]
        cp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r3 = cp2.add_run(f"Guide Signature :\n{w['guide_sign']}")
        r3.bold = True
        r3.font.size = Pt(8.5)
        r3.font.color.rgb = c_blue

        doc.add_page_break()

    # Reorder XML schemas for 100% MS Word compliance
    reorder_children(doc.sections[0]._sectPr, SECTPR_ORDER)
    for t in doc.tables:
        reorder_children(t._tbl.tblPr, TBLPR_ORDER)
        for row in t.rows:
            for cell in row.cells:
                tcPr = cell._tc.find(qn('w:tcPr'))
                if tcPr is not None:
                    reorder_children(tcPr, TCPR_ORDER)

    doc.save(docx_path)
    print(f"[SUCCESS] Built {docx_path} successfully!")

if __name__ == '__main__':
    build_docx_workbook()
